"""
from pptx import Presentation
from pptx.util import Inches

# Create a new PowerPoint presentation
prs = Presentation()

# Define slide titles and descriptions (used for speaker notes)
slides_info = [
    ("Tech to Prioritize to Remain Competitive", "Introduction slide with your name and date."),
    ("The Need for Tech Prioritization", "Today’s market moves fast. Prioritizing tech is essential to stay ahead."),
    ("AI & Automation", "AI improves efficiency and reduces costs, from resume screening to predictive analytics."),
    ("Cloud Computing & Scalability", "Cloud platforms let businesses scale and adapt quickly."),
    ("Cybersecurity Investments", "With rising cyber threats, proactive defense is non-negotiable."),
    ("Data-Driven Decision Making", "Smart decisions are made with reliable, real-time data."),
    ("Remote Collaboration Tools", "Work from anywhere is here to stay — tools must support it."),
    ("Talent & Tech Together", "Tech without the right talent doesn’t deliver results."),
    ("Risks of Falling Behind", "Failure to prioritize means lost market share and talent."),
    ("Final Takeaway", "Invest Smart. Prioritize Right. Be the company top talent wants to join."),
]

def add_image_slide(prs, title, notes, image_file=None):
    slide_layout = prs.slide_layouts[5]  # Title Only layout
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title
    
    notes_slide = slide.notes_slide
    notes_slide.notes_text_frame.text = notes
    
    if image_file:
        left = Inches(1)
        top = Inches(1.5)
        height = Inches(5)
        slide.shapes.add_picture(image_file, left, top, height=height)


# Loop through slides and add them
for i, (title, notes) in enumerate(slides_info, start=1):
    image_file = f"image{i}.jpg"  # Change to your own path or files
    add_image_slide(prs, title, notes, image_file=image_file)

# Save the presentation
prs.save("Tech_to_Prioritize_Presentation.pptx")
print("Presentation successfully saved.")
